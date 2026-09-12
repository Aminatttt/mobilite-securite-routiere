from pathlib import Path
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# ======================================
# CHEMINS
# ======================================

RACINE = Path(__file__).resolve().parent.parent

FICHIER_ENTREE = (
    RACINE
    / "data"
    / "curated"
    / "dictionnaire_donnees_enrichi.csv"
)

FICHIER_SORTIE = (
    RACINE
    / "data"
    / "curated"
    / "dictionnaire_donnees_partage.xlsx"
)


# ======================================
# LECTURE DU FICHIER ENRICHI
# ======================================

print("======================================")
print("CREATION DU DICTIONNAIRE EXCEL")
print("======================================")

print(f"\nLecture : {FICHIER_ENTREE}")

df = pd.read_csv(
    FICHIER_ENTREE,
    encoding="utf-8-sig"
)

print(f"Variables trouvées : {len(df)}")


# ======================================
# RENOMMAGE DES COLONNES
# ======================================

df = df.rename(columns={

    "source":
        "Source",

    "fichier":
        "Fichier",

    "variable":
        "Variable",

    "type":
        "Type",

    "nb_valeurs":
        "Nb valeurs",

    "nb_manquantes":
        "Nb manquantes",

    "pct_manquantes":
        "% manquantes",

    "exemple":
        "Exemple",

    "cle_jointure":
        "Clé de jointure",

    "unite":
        "Unité",

    "description":
        "Description"
})


# ======================================
# CREATION DU FICHIER EXCEL
# ======================================

df.to_excel(
    FICHIER_SORTIE,
    index=False,
    sheet_name="Dictionnaire"
)

print("\nFichier Excel créé.")


# ======================================
# OUVERTURE DU FICHIER EXCEL
# ======================================

wb = load_workbook(FICHIER_SORTIE)

ws = wb["Dictionnaire"]


# ======================================
# STYLE DE L'EN-TETE
# ======================================

couleur_entete = "1F4E78"

for cellule in ws[1]:

    cellule.fill = PatternFill(
        fill_type="solid",
        fgColor=couleur_entete
    )

    cellule.font = Font(
        bold=True,
        color="FFFFFF"
    )

    cellule.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )


# Hauteur de l'en-tête

ws.row_dimensions[1].height = 30


# ======================================
# CREATION DU TABLEAU EXCEL
# ======================================

derniere_ligne = ws.max_row

derniere_colonne = ws.max_column

ref_tableau = (
    f"A1:{get_column_letter(derniere_colonne)}"
    f"{derniere_ligne}"
)

tableau = Table(
    displayName="DictionnaireDonnees",
    ref=ref_tableau
)

style_tableau = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False
)

tableau.tableStyleInfo = style_tableau

ws.add_table(tableau)


# ======================================
# FIGER LA PREMIERE LIGNE
# ======================================

ws.freeze_panes = "A2"


# ======================================
# ALIGNEMENT ET RETOUR A LA LIGNE
# ======================================

for ligne in ws.iter_rows():

    for cellule in ligne:

        cellule.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )


# ======================================
# LARGEUR DES COLONNES
# ======================================

largeurs = {

    "A": 18,   # Source

    "B": 42,   # Fichier

    "C": 28,   # Variable

    "D": 15,   # Type

    "E": 14,   # Nb valeurs

    "F": 16,   # Nb manquantes

    "G": 16,   # % manquantes

    "H": 25,   # Exemple

    "I": 35,   # Clé de jointure

    "J": 20,   # Unité

    "K": 60    # Description
}


for colonne, largeur in largeurs.items():

    ws.column_dimensions[colonne].width = largeur


# ======================================
# MISE EN EVIDENCE DES CLES DE JOINTURE
# ======================================

colonne_cle = None

for cellule in ws[1]:

    if cellule.value == "Clé de jointure":

        colonne_cle = cellule.column

        break


if colonne_cle:

    for ligne in range(2, ws.max_row + 1):

        cellule = ws.cell(
            row=ligne,
            column=colonne_cle
        )

        valeur = str(
            cellule.value
        ).strip().upper()

        if valeur.startswith("OUI"):

            cellule.font = Font(
                bold=True
            )


# ======================================
# BORDURES
# ======================================

bordure = Border(
    bottom=Side(
        style="thin",
        color="D9E1F2"
    )
)


for ligne in ws.iter_rows():

    for cellule in ligne:

        cellule.border = bordure


# ======================================
# FORMAT DES POURCENTAGES
# ======================================

colonne_pourcentage = None

for cellule in ws[1]:

    if cellule.value == "% manquantes":

        colonne_pourcentage = cellule.column

        break


if colonne_pourcentage:

    for ligne in range(2, ws.max_row + 1):

        cellule = ws.cell(
            row=ligne,
            column=colonne_pourcentage
        )

        if isinstance(cellule.value, (int, float)):

            cellule.number_format = "0.00"


# ======================================
# SAUVEGARDE FINALE
# ======================================

wb.save(FICHIER_SORTIE)


# ======================================
# MESSAGE FINAL
# ======================================

print("\n======================================")
print("DICTIONNAIRE EXCEL CREE")
print("======================================")

print(f"Fichier : {FICHIER_SORTIE}")

print(f"Variables : {len(df)}")

print("\nFonctionnalités ajoutées :")

print("- Tableau Excel")
print("- Filtres")
print("- En-tête coloré")
print("- Lignes alternées")
print("- Première ligne figée")
print("- Largeurs de colonnes")
print("- Retour automatique à la ligne")
print("- Mise en évidence des clés de jointure")
print("- Format des pourcentages")

print("\n======================================")
print("✅ TERMINÉ")
print("======================================")